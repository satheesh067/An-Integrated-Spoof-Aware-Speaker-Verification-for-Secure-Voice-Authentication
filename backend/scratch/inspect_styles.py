from pptx import Presentation
from pptx.dml.color import RGBColor

prs = Presentation("Secure-Biometric-Voice-Authentication-Using-NLP-Based-Speaker-Verification-and-Anti-Spoofing.pptx")
print("Slide Width:", prs.slide_width, "Slide Height:", prs.slide_height)

for idx, slide in enumerate(prs.slides):
    print(f"\n--- Slide {idx+1} ---")
    for shape in slide.shapes:
        if shape.has_text_frame:
            for p in shape.text_frame.paragraphs:
                for run in p.runs:
                    font = run.font
                    color = font.color.rgb if font.color and font.color.type == 1 else "default"
                    print(f"[{shape.name}] Text: {run.text[:30]!r} | Font: {font.name} | Size: {font.size} | Color: {color} | Bold: {font.bold}")
                    break
                if p.runs:
                    break
