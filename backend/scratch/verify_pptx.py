from pptx import Presentation

prs = Presentation("Secure-Biometric-Voice-Authentication-Using-NLP-Based-Speaker-Verification-and-Anti-Spoofing.pptx")

with open("scratch/verify_utf8.txt", "w", encoding="utf-8") as f:
    f.write(f"Total slides in updated presentation: {len(prs.slides)}\n")
    for idx, slide in enumerate(prs.slides):
        f.write(f"\n==================== Slide {idx + 1} ====================\n")
        for s_idx, shape in enumerate(slide.shapes):
            if shape.has_text_frame:
                texts = [p.text.strip() for p in shape.text_frame.paragraphs if p.text.strip()]
                if texts:
                    f.write(f"[{shape.name}] (type={shape.shape_type}):\n")
                    for t in texts:
                        f.write(f"    - {t}\n")
            elif shape.has_table:
                f.write(f"[{shape.name} TABLE]: {len(shape.table.rows)} rows\n")
            elif shape.shape_type == 13: # Picture
                f.write(f"[{shape.name} PICTURE]: pos=({shape.left}, {shape.top}) size=({shape.width}, {shape.height})\n")

print("Saved scratch/verify_utf8.txt")
