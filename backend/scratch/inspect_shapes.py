import sys
from pathlib import Path
from pptx import Presentation

pptx_path = Path("Secure-Biometric-Voice-Authentication-Using-NLP-Based-Speaker-Verification-and-Anti-Spoofing.pptx")
prs = Presentation(pptx_path)

print(f"Total slides: {len(prs.slides)}")
for idx, slide in enumerate(prs.slides):
    print(f"\n--- Slide {idx + 1} (Layout: {slide.slide_layout.name}) ---")
    for s_idx, shape in enumerate(slide.shapes):
        shape_type = shape.shape_type
        print(f"  Shape {s_idx}: {shape.name} (type={shape_type}, left={shape.left}, top={shape.top}, width={shape.width}, height={shape.height})")
        if shape.has_text_frame:
            for p in shape.text_frame.paragraphs:
                if p.text.strip():
                    print(f"    Text: {p.text.strip()}")
        elif shape.has_table:
            print("    [Table]")
