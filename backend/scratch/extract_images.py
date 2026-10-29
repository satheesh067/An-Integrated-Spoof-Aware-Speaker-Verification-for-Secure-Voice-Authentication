import os
from pathlib import Path
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

pptx_path = Path("Secure-Biometric-Voice-Authentication-Using-NLP-Based-Speaker-Verification-and-Anti-Spoofing.pptx")
prs = Presentation(pptx_path)

out_dir = Path("scratch/extracted_images")
out_dir.mkdir(exist_ok=True)

for s_idx, slide in enumerate(prs.slides):
    for sh_idx, shape in enumerate(slide.shapes):
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            img = shape.image
            filename = f"slide_{s_idx+1}_img_{sh_idx}_{shape.name}.{img.ext}"
            filepath = out_dir / filename
            with open(filepath, "wb") as f:
                f.write(img.blob)
            print(f"Saved {filepath} ({len(img.blob)} bytes)")
