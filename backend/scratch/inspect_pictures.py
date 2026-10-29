import sys
from pathlib import Path
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

pptx_path = Path("Secure-Biometric-Voice-Authentication-Using-NLP-Based-Speaker-Verification-and-Anti-Spoofing.pptx")
prs = Presentation(pptx_path)

with open("scratch/shapes_dump.txt", "w", encoding="utf-8") as f:
    for idx, slide in enumerate(prs.slides):
        f.write(f"\n--- Slide {idx + 1} ---\n")
        for s_idx, shape in enumerate(slide.shapes):
            f.write(f"Shape {s_idx}: name={shape.name}, type={shape.shape_type}\n")
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                image = shape.image
                f.write(f"  [PICTURE] format={image.ext}, size={len(image.blob)} bytes\n")
            if shape.has_text_frame:
                for p in shape.text_frame.paragraphs:
                    if p.text.strip():
                        f.write(f"  [TEXT] {p.text.strip()}\n")

print("Dumped shapes to scratch/shapes_dump.txt")
