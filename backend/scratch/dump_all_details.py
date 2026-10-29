import sys
from pathlib import Path
from pptx import Presentation

prs = Presentation("Secure-Biometric-Voice-Authentication-Using-NLP-Based-Speaker-Verification-and-Anti-Spoofing.pptx")

with open("scratch/all_shapes_detailed.txt", "w", encoding="utf-8") as f:
    for idx, slide in enumerate(prs.slides):
        f.write(f"\n==================== Slide {idx+1} ====================\n")
        for s_idx, shape in enumerate(slide.shapes):
            f.write(f"\nShape {s_idx}: name={shape.name}, type={shape.shape_type}, left={shape.left}, top={shape.top}, width={shape.width}, height={shape.height}\n")
            if shape.has_text_frame:
                tf = shape.text_frame
                for p_idx, p in enumerate(tf.paragraphs):
                    runs_txt = "".join(r.text for r in p.runs)
                    f.write(f"  P{p_idx} (level={p.level}): {runs_txt}\n")
            elif shape.has_table:
                f.write(f"  [TABLE]: {len(shape.table.rows)} rows, {len(shape.table.columns)} cols\n")

print("Dumped all shapes detailed.")
