from pptx import Presentation

prs = Presentation("Secure-Biometric-Voice-Authentication-Using-NLP-Based-Speaker-Verification-and-Anti-Spoofing.pptx")

for idx, slide in enumerate(prs.slides):
    print(f"\n==================== Slide {idx + 1} ====================")
    for s_idx, shape in enumerate(slide.shapes):
        print(f"Shape {s_idx}: name={shape.name}, type={shape.shape_type}, left={shape.left}, top={shape.top}, width={shape.width}, height={shape.height}")
        if shape.has_text_frame:
            txt = shape.text.replace('\n', ' ')[:60]
            print(f"   Text: {txt!r}")
