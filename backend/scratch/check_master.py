from pptx import Presentation

prs = Presentation("Secure-Biometric-Voice-Authentication-Using-NLP-Based-Speaker-Verification-and-Anti-Spoofing.pptx")
slide = prs.slides[0]
print("Slide master count:", len(prs.slide_masters))
master = prs.slide_masters[0]
print("Master shapes:", len(master.shapes))
for s in master.shapes:
    print(" Master shape:", s.name, s.shape_type)
    if hasattr(s, "fill"):
        print("  fill type:", s.fill.type)

layout = slide.slide_layout
print("Layout shapes:", len(layout.shapes))
for s in layout.shapes:
    print(" Layout shape:", s.name, s.shape_type)
