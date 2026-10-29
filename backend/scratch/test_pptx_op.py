import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

prs = Presentation("Secure-Biometric-Voice-Authentication-Using-NLP-Based-Speaker-Verification-and-Anti-Spoofing.pptx.backup")
slide = prs.slides[8] # Slide 9 (0-indexed 8)
print("Slide 9 shape count:", len(slide.shapes))
for s in slide.shapes:
    print(s.name, s.shape_type)

print("Test load succeeded.")
