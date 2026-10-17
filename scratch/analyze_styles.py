from pptx import Presentation

prs = Presentation("Secure-Biometric-Voice-Authentication-Using-NLP-Based-Speaker-Verification-and-Anti-Spoofing.pptx")
with open("scratch/style_analysis.txt", "w", encoding="utf-8") as f:
    for idx, slide in enumerate(prs.slides):
        f.write(f"\n==================== Slide {idx+1} ====================\n")
        bg = slide.background
        f.write(f"Background fill type: {bg.fill.type if bg else 'None'}\n")
        for s_idx, shape in enumerate(slide.shapes):
            f.write(f"Shape {s_idx}: {shape.name} (type={shape.shape_type})\n")
            if shape.has_text_frame:
                for p_idx, p in enumerate(shape.text_frame.paragraphs):
                    for r_idx, run in enumerate(p.runs):
                        f.write(f"  P{p_idx} R{r_idx}: {run.text[:40]!r} | font={run.font.name} | size={run.font.size} | bold={run.font.bold}\n")

print("Saved style analysis to scratch/style_analysis.txt")
