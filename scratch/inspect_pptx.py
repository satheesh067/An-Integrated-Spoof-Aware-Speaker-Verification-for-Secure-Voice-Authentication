import sys
from pathlib import Path
from pptx import Presentation

pptx_path = Path("Secure-Biometric-Voice-Authentication-Using-NLP-Based-Speaker-Verification-and-Anti-Spoofing.pptx")
if not pptx_path.exists():
    print("File does not exist:", pptx_path)
    sys.exit(1)

prs = Presentation(pptx_path)
out_path = Path("scratch/pptx_dump.txt")

with open(out_path, "w", encoding="utf-8") as f:
    f.write(f"Slide width: {prs.slide_width}, Slide height: {prs.slide_height}\n")
    f.write(f"Total slides: {len(prs.slides)}\n\n")

    for idx, slide in enumerate(prs.slides):
        f.write(f"\n==================== Slide {idx + 1} ====================\n")
        for shape in slide.shapes:
            if shape.has_text_frame:
                text = shape.text.strip()
                if text:
                    f.write(f"\n[{shape.name}] (left={shape.left}, top={shape.top}, width={shape.width}, height={shape.height}):\n")
                    for p_idx, p in enumerate(shape.text_frame.paragraphs):
                        f.write(f"  P{p_idx} (level={p.level}): {p.text}\n")
            elif shape.has_table:
                f.write(f"\n[TABLE in {shape.name}]:\n")
                for row_idx, row in enumerate(shape.table.rows):
                    row_text = [cell.text.replace('\n', ' ') for cell in row.cells]
                    f.write(f"  Row {row_idx}: {' | '.join(row_text)}\n")

print("Dumped to scratch/pptx_dump.txt successfully.")
