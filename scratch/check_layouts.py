from pptx import Presentation

prs = Presentation("Secure-Biometric-Voice-Authentication-Using-NLP-Based-Speaker-Verification-and-Anti-Spoofing.pptx")

for idx, slide in enumerate(prs.slides):
    print(f"Slide {idx+1}: layout={slide.slide_layout.name}, shapes={len(slide.shapes)}")
